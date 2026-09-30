import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import io
import os
import re
import base64
import pdfplumber
import docx
from datetime import datetime

# ReportLab para geração de PDF
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch

# =========================================================
# CONFIGURAÇÃO DA PÁGINA E TEMA VISUAL (DARK MODE)
# =========================================================
st.set_page_config(page_title="Portal Gerador de Relatórios — PM/3", page_icon="🛡️", layout="centered")

def obter_caminho_brasao():
    for nome in ["brasao.png", "brasao.PNG", "Brasao.png", "BRASAO.PNG", "brasao.jpg", "brasao.jpeg"]:
        if os.path.exists(nome):
            return nome
    return None

caminho_brasao = obter_caminho_brasao()

def gerar_css_app(caminho_img):
    css_base = """
    <style>
        div[data-testid="stFileUploader"], div[data-testid="stTextInput"], div[data-testid="stTextArea"] {
            background-color: rgba(30, 34, 45, 0.88) !important;
            border-radius: 10px;
            padding: 10px;
            border: 1px solid #2E364A;
        }
        .portal-card {
            background: linear-gradient(135deg, rgba(30, 38, 56, 0.95), rgba(18, 23, 35, 0.95));
            border: 1px solid #2E364A;
            border-radius: 12px;
            padding: 20px;
            text-align: center;
            box-shadow: 0 4px 12px rgba(0,0,0,0.4);
            margin-bottom: 15px;
        }
        .portal-card h3 {
            color: #60A5FA !important;
            margin-bottom: 8px;
        }
        .portal-card p {
            color: #9CA3AF;
            font-size: 13.5px;
            margin-bottom: 12px;
        }
        .stButton>button {
            background-color: #002060;
            color: #FFFFFF;
            font-weight: bold;
            border-radius: 8px;
            border: 1px solid #1E3A8A;
            width: 100%;
            padding: 10px;
        }
        .stButton>button:hover {
            background-color: #1E40AF;
            border-color: #3B82F6;
        }
        h1, h2, h3, h4, h5 {
            color: #F3F4F6 !important;
        }
    </style>
    """
    
    if caminho_img and os.path.exists(caminho_img):
        try:
            with open(caminho_img, "rb") as f:
                encoded = base64.b64encode(f.read()).decode("utf-8")
            ext = caminho_img.split(".")[-1].lower()
            mime = "image/png" if "png" in ext else "image/jpeg"
            bg_css = f"""
            <style>
                .stApp {{
                    background-color: #0E1117;
                    color: #E0E6ED;
                    background-image: linear-gradient(rgba(14, 17, 23, 0.92), rgba(14, 17, 23, 0.92)), url('data:{mime};base64,{encoded}');
                    background-size: contain;
                    background-repeat: no-repeat;
                    background-position: center center;
                    background-attachment: fixed;
                }}
            </style>
            """
            return bg_css + css_base
        except Exception:
            pass

    return """
    <style>
        .stApp {
            background-color: #0E1117;
            color: #E0E6ED;
        }
    </style>
    """ + css_base

st.markdown(gerar_css_app(caminho_brasao), unsafe_allow_html=True)

# 🏷 ASSINATURA MOVIDA PARA A DIREITA
st.markdown("""
<div style="position: fixed; bottom: 15px; right: 140px; text-align: right; color: #9CA3AF; font-size: 12px; font-family: sans-serif; z-index: 999999; line-height: 1.4; background-color: rgba(14, 17, 23, 0.9); padding: 6px 12px; border-radius: 6px; border: 1px solid #2E364A;">
    Desenvolvido por:<br>
    <strong style="color: #60A5FA; font-size: 13px;">Nathan Wenzel</strong>
</div>
""", unsafe_allow_html=True)

# =========================================================
# CONTROLE DE SESSÃO E AUTENTICAÇÃO
# =========================================================
SENHA_CORRETA = "deusa"

if "autenticado" not in st.session_state:
    st.session_state.autenticado = False

if "pagina_atual" not in st.session_state:
    st.session_state.pagina_atual = "portal"

if not st.session_state.autenticado:
    if caminho_brasao:
        col1, col2, col3 = st.columns(3)
        with col2:
            st.image(caminho_brasao, width=190)

    st.title("🔒 Acesso Restrito — 18º BPM")
    st.write("Digite a senha de acesso para utilizar o Portal de Sistemas Operacionais.")
    
    senha_input = st.text_input("Senha de acesso:", type="password")
    
    if st.button("Entrar no Portal"):
        if senha_input == SENHA_CORRETA:
            st.session_state.autenticado = True
            st.success("Acesso liberado!")
            st.rerun()
        else:
            st.error("Senha incorreta! Verifique e tente novamente.")
    st.stop()


# =========================================================
# FUNÇÕES AUXILIARES — LEITURA E PROCESSAMENTO
# =========================================================

MESES = {
    1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril",
    5: "maio", 6: "junho", 7: "julho", 8: "agosto",
    9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro"
}

def data_atual_extenso():
    now = datetime.now()
    return f"{now.day} de {MESES[now.month]} de {now.year}"

def set_cell_bg(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def limpar_assinaturas_e_ruidos(texto):
    if not texto:
        return ""
    padroes_lixo = [
        r'.*Assinado\s+eletronicamente.*',
        r'Documento\s+assinado\s+digitalmente.*',
        r'Inserido\s+ao\s+protocolo.*',
        r'conforme\s+MP\s+n[º°\.]?\s*2\.?200-2/2001.*',
        r'Validação:.*',
        r'https?://[^\s]*eprotocolo[^\s]*',
        r'www\.[^\s]*eprotocolo[^\s]*',
        r'Código\s+de\s+autenticidade:.*',
        r'e-Protocolo\s*\d+.*',
        r'SHA-?256:.*',
        r'Página\s+\d+\s+de\s+\d+.*',
    ]
    linhas = texto.split('\n')
    linhas_limpas = []
    for linha in linhas:
        l_trim = linha.strip()
        if not l_trim:
            continue
        descartar = False
        for p in padroes_lixo:
            if re.search(p, l_trim, re.IGNORECASE):
                descartar = True
                break
        if not descartar:
            linhas_limpas.append(l_trim)
    return '\n'.join(linhas_limpas).strip()

def safe_crop_text(page, y0, y1):
    if y1 <= y0 + 1:
        return ""
    try:
        cropped = page.crop((0, y0, page.width, y1))
        return cropped.extract_text() or ""
    except Exception:
        return ""

def is_real_data_table(extracted_tbl):
    if not extracted_tbl:
        return False
    max_cols = max(len([c for c in r if c and str(c).strip()]) for r in extracted_tbl if r)
    return max_cols >= 2

def extrair_texto_arquivo(uploaded_file):
    ext = uploaded_file.name.split(".")[-1].lower()
    text = ""
    
    if ext == "pdf":
        uploaded_file.seek(0)
        with pdfplumber.open(uploaded_file) as pdf:
            for page in pdf.pages:
                tables = page.find_tables()
                if not tables:
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        text += page_text + "\n\n"
                else:
                    for t in tables:
                        extracted_tbl = t.extract()
                        if extracted_tbl:
                            table_lines = []
                            for row in extracted_tbl:
                                if row and any(c is not None and str(c).strip() != "" for c in row):
                                    clean_row = [str(c).replace('\n', ' ').strip() if c is not None else "" for c in row]
                                    table_lines.append(" | ".join(clean_row))
                            if table_lines:
                                text += "\n".join(table_lines) + "\n\n"
        uploaded_file.seek(0)
    elif ext in ["docx", "doc"]:
        uploaded_file.seek(0)
        doc = docx.Document(uploaded_file)
        items = []
        for element in doc.element.body:
            if element.tag.endswith('p'):
                p = docx.text.paragraph.Paragraph(element, doc)
                if p.text.strip():
                    items.append(p.text.strip())
            elif element.tag.endswith('tbl'):
                tbl = docx.table.Table(element, doc)
                table_matrix = []
                for row in tbl.rows:
                    row_cells = [c.text.replace('\n', ' ').strip() for c in row.cells]
                    if any(cell for cell in row_cells):
                        table_matrix.append(row_cells)
                if table_matrix:
                    table_lines = [" | ".join(r) for r in table_matrix]
                    items.append("\n".join(table_lines))
        text = "\n\n".join(items)
        uploaded_file.seek(0)
        
    return limpar_assinaturas_e_ruidos(text)


# =========================================================
# GERADORES DE DOCX E PDF — EXTRAJORNADA
# =========================================================

def gerar_relatorio_extrajornada_docx(titulo_relatorio, data_servico, comandante_nome, cargo_funcao, conteudo_tabela):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Cabeçalho
    p_hdr = doc.add_paragraph()
    p_hdr.paragraph_format.line_spacing = 1.15
    p_hdr.paragraph_format.space_after = Pt(12)
    r_hdr = p_hdr.add_run("POLÍCIA MILITAR DO PARANÁ\n2º COMANDO REGIONAL DE POLÍCIA MILITAR\n18º BATALHÃO DE POLÍCIA MILITAR — P/3")
    r_hdr.bold = True
    r_hdr.font.name = "Arial"
    r_hdr.font.size = Pt(10)

    # Título
    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_before = Pt(12)
    p_tit.paragraph_format.space_after = Pt(14)
    r_tit = p_tit.add_run(titulo_relatorio.upper())
    r_tit.bold = True
    r_tit.font.name = "Arial"
    r_tit.font.size = Pt(12)
    r_tit.font.color.rgb = RGBColor(0, 32, 96)

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_after = Pt(10)
    r_info = p_info.add_run(f"Data de Referência/Emprego: {data_servico}\nLocal de Atuação: Circunscrição do 18º BPM")
    r_info.font.name = "Arial"
    r_info.font.size = Pt(10)

    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(10)
    p_sec.paragraph_format.space_after = Pt(6)
    r_sec = p_sec.add_run("1. PROGRAMAÇÃO E ESCALA DE SERVIÇO VOLUNTÁRIO")
    r_sec.bold = True
    r_sec.font.name = "Arial"
    r_sec.font.size = Pt(11)

    # Processar linhas da tabela
    linhas = conteudo_tabela.strip().split('\n')
    tabela_matrix = []
    for l in linhas:
        if '|' in l:
            cels = [c.strip() for c in l.split('|') if c.strip()]
            if cels:
                tabela_matrix.append(cels)

    if tabela_matrix:
        max_cols = max(len(r) for r in tabela_matrix)
        tbl = doc.add_table(rows=len(tabela_matrix), cols=max_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.style = 'Table Grid'
        
        for r_idx, r_data in enumerate(tabela_matrix):
            for c_idx, val in enumerate(r_data):
                if c_idx < max_cols:
                    cell = tbl.rows[r_idx].cells[c_idx]
                    p = cell.paragraphs
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    r = p.add_run(val)
                    r.font.name = "Arial"
                    if r_idx == 0:
                        set_cell_bg(cell, "002060")
                        r.bold = True
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        r.font.size = Pt(8.5)

    # Assinatura
    p_ass = doc.add_paragraph()
    p_ass.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ass.paragraph_format.space_before = Pt(30)
    r_a1 = p_ass.add_run(f"{comandante_nome}\n")
    r_a1.bold = True
    r_a1.font.name = "Arial"
    r_a1.font.size = Pt(10)
    r_a2 = p_ass.add_run(cargo_funcao)
    r_a2.font.name = "Arial"
    r_a2.font.size = Pt(9.5)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

def gerar_relatorio_extrajornada_pdf(titulo_relatorio, data_servico, comandante_nome, cargo_funcao, conteudo_tabela):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=0.75*inch, leftMargin=0.75*inch, topMargin=0.75*inch, bottomMargin=0.75*inch)
    styles = getSampleStyleSheet()
    
    style_hdr = ParagraphStyle('Hdr', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, leading=13)
    style_tit = ParagraphStyle('Tit', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=12, leading=16, alignment=1, textColor=colors.HexColor('#002060'))
    style_body = ParagraphStyle('Body', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=12)
    style_th = ParagraphStyle('TH', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8.5, leading=11, textColor=colors.white)
    style_ass = ParagraphStyle('Ass', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=13, alignment=1)

    story = []
    story.append(Paragraph("POLÍCIA MILITAR DO PARANÁ<br/>2º COMANDO REGIONAL DE POLÍCIA MILITAR<br/>18º BATALHÃO DE POLÍCIA MILITAR — P/3", style_hdr))
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=2, spaceAfter=10))
    story.append(Paragraph(f"<b>{titulo_relatorio.upper()}</b>", style_tit))
    story.append(Spacer(1, 10))
    story.append(Paragraph(f"Data de Referência: {data_servico}<br/>Local de Atuação: Circunscrição do 18º BPM", style_body))
    story.append(Spacer(1, 10))

    linhas = conteudo_tabela.strip().split('\n')
    tabela_matrix = []
    for l in linhas:
        if '|' in l:
            cels = [c.strip() for c in l.split('|') if c.strip()]
            if cels:
                tabela_matrix.append(cels)

    if tabela_matrix:
        max_cols = max(len(r) for r in tabela_matrix)
        col_w = (6.4 * inch) / max_cols
        pdf_data = []
        for r_idx, r_data in enumerate(tabela_matrix):
            row_p = []
            for c_val in r_data:
                c_clean = c_val.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                if r_idx == 0:
                    row_p.append(Paragraph(c_clean, style_th))
                else:
                    row_p.append(Paragraph(c_clean, style_body))
            while len(row_p) < max_cols:
                row_p.append(Paragraph("", style_body))
            pdf_data.append(row_p)
            
        tbl = Table(pdf_data, colWidths=[col_w]*max_cols)
        tbl.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#002060')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#A0AAB5')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')])
        ]))
        story.append(tbl)

    story.append(Spacer(1, 25))
    story.append(Paragraph(f"<b>{comandante_nome}</b><br/>{cargo_funcao}", style_ass))
    doc.build(story)
    buf.seek(0)
    return buf.getvalue()


# =========================================================
# MÓDULO 1: GERADOR DE OS
# =========================================================
def render_modulo_gerador_os():
    st.title("📑 Gerador de Ordem de Serviço (OS)")
    st.caption("18º Batalhão de Polícia Militar — PMPR")
    st.write("Envie a **Ordem de Operação (OO)** para extração automática e geração da **Ordem de Serviço (OS)**.")

    arquivo_oo = st.file_uploader("Envie o arquivo da Ordem de Operação (PDF ou DOCX)", type=["pdf", "docx", "doc"], key="uploader_os")

    if arquivo_oo:
        texto_extraido = extrair_texto_arquivo(arquivo_oo)
        if texto_extraido.strip():
            st.success("Ordem de Operação analisada com sucesso!")
            
            st.subheader("📝 Informações da Ordem de Serviço")
            col_a, col_b = st.columns(2)
            with col_a:
                num_os = st.text_input("Número da OS", value="077", key="os_num")
                data_expedicao = st.text_input("Data de Expedição", value=data_atual_extenso(), key="os_data")
            with col_b:
                nome_operacao = st.text_input("Nome da Operação", value="OPERAÇÃO POLICIAL", key="os_nome")
                nome_comandante = st.text_input("Comandante", value="Ten.-Cel. QOEM PM Helder de Lima Dantas Junior", key="os_cmd")

            if st.button("🚀 Gerar Ordem de Serviço (Word e PDF)", key="btn_gerar_os"):
                fields_final = {
                    'num_os': num_os,
                    'data_expedicao': data_expedicao,
                    'nome_operacao': nome_operacao,
                    'finalidade': "Realizar ações de policiamento ostensivo e preservação da ordem pública.",
                    'informacoes_gerais': texto_extraido[:800],
                    'missao': "Atuação ostensiva e preventiva na circunscrição do 18º BPM.",
                    'execucao': "Emprego de equipes operacionais.",
                    'logistica': "Equipamentos e armamentos orgânicos da OPM.",
                    'relatorios': "Lançamento no SISGCOP.",
                    'prescricoes': "Estrito cumprimento do dever legal.",
                    'referencias': "Constituição Federal; LOB PMPR.",
                    'nome_comandante': nome_comandante,
                    'cargo_comandante': "Comandante do 18º BPM."
                }
                docx_b = gerar_relatorio_extrajornada_docx("ORDEM DE SERVIÇO Nº " + num_os, data_expedicao, nome_comandante, "Comandante do 18º BPM.", texto_extraido[:1000])
                pdf_b = gerar_relatorio_extrajornada_pdf("ORDEM DE SERVIÇO Nº " + num_os, data_expedicao, nome_comandante, "Comandante do 18º BPM.", texto_extraido[:1000])
                
                st.session_state['os_docx'] = docx_b
                st.session_state['os_pdf'] = pdf_b

            if 'os_docx' in st.session_state:
                st.markdown("---")
                col_d1, col_d2 = st.columns(2)
                with col_d1:
                    st.download_button("📄 Baixar OS em Word (.docx)", st.session_state['os_docx'], file_name=f"OS_{num_os}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                with col_d2:
                    st.download_button("📕 Baixar OS em PDF (.pdf)", st.session_state['os_pdf'], file_name=f"OS_{num_os}.pdf", mime="application/pdf")


# =========================================================
# MÓDULO 2: GERADOR DE RELATÓRIO EXTRAJORNADA (COM UPLOAD E DOWNLOAD)
# =========================================================
def render_modulo_segundo_site():
    st.title("🗓️ Gerador de Relatório de Extrajornada")
    st.caption("18º Batalhão de Polícia Militar — PMPR")
    st.write("Envie a **Escala / Programação Extrajornada** (PDF, DOCX, XLSX ou CSV) para processar e gerar o **Relatório Formatado** para download.")

    arquivo_escala = st.file_uploader("Envie o arquivo de Escala / Programação Extrajornada", type=["pdf", "docx", "doc", "xlsx", "csv"], key="uploader_extra")

    if arquivo_escala:
        st.success(f"Arquivo '{arquivo_escala.name}' carregado com sucesso!")
        
        ext = arquivo_escala.name.split(".")[-1].lower()
        conteudo_tabela_extraida = ""
        
        if ext in ["pdf", "docx", "doc"]:
            conteudo_tabela_extraida = extrair_texto_arquivo(arquivo_escala)
        elif ext in ["xlsx", "csv"]:
            try:
                df = pd.read_excel(arquivo_escala) if ext == "xlsx" else pd.read_csv(arquivo_escala)
                st.dataframe(df, use_container_width=True)
                # Converte dataframe para linhas formatadas
                conteudo_tabela_extraida = " | ".join([str(c) for c in df.columns]) + "\n"
                for _, row in df.iterrows():
                    conteudo_tabela_extraida += " | ".join([str(v) if pd.notna(v) else "" for v in row.values]) + "\n"
            except Exception as e:
                st.error(f"Erro ao ler a planilha: {e}")

        st.subheader("📝 Configuração do Relatório de Extrajornada")
        
        col1, col2 = st.columns(2)
        with col1:
            titulo_relatorio = st.text_input("Título do Relatório / Programação", value="RELATÓRIO DE PROGRAMAÇÃO EXTRAJORNADA VOLUNTÁRIA", key="extra_tit")
            data_referencia = st.text_input("Data de Referência", value=data_atual_extenso(), key="extra_data")
        with col2:
            comandante_nome = st.text_input("Comandante / Responsável", value="Ten.-Cel. QOEM PM Helder de Lima Dantas Junior", key="extra_cmd")
            cargo_funcao = st.text_input("Cargo / Função", value="Comandante do 18º BPM", key="extra_cargo")

        st.subheader("📊 Tabela / Dados Extraídos da Escala")
        tabela_editavel = st.text_area("Ajuste os dados da escala abaixo se necessário:", value=conteudo_tabela_extraida, height=250, key="extra_txt")

        if st.button("🚀 Gerar Relatório de Extrajornada (Word e PDF)", key="btn_gerar_extra"):
            docx_extra = gerar_relatorio_extrajornada_docx(titulo_relatorio, data_referencia, comandante_nome, cargo_funcao, tabela_editavel)
            pdf_extra = gerar_relatorio_extrajornada_pdf(titulo_relatorio, data_referencia, comandante_nome, cargo_funcao, tabela_editavel)
            
            st.session_state['extra_docx'] = docx_extra
            st.session_state['extra_pdf'] = pdf_extra

        if 'extra_docx' in st.session_state:
            st.markdown("---")
            st.subheader("📥 Baixar Relatório Gerado")
            col_ex1, col_d2 = st.columns(2)
            with col_ex1:
                st.download_button("📄 Baixar Relatório em Word (.docx)", st.session_state['extra_docx'], file_name="Relatorio_Extrajornada_18BPM.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            with col_d2:
                st.download_button("📕 Baixar Relatório em PDF (.pdf)", st.session_state['extra_pdf'], file_name="Relatorio_Extrajornada_18BPM.pdf", mime="application/pdf")


# =========================================================
# TELA DE ENTRADA (PORTAL / HUB CENTRAL)
# =========================================================
def render_tela_entrada():
    if caminho_brasao:
        c1, c2, c3 = st.columns(3)
        with c2:
            st.image(caminho_brasao, width=220)

    st.markdown("""
    <div style='text-align: center; margin-bottom: 25px; margin-top: 10px;'>
        <h1 style='font-size: 26px; font-weight: bold; margin-bottom: 6px;'>🛡️ Portal Gerador de Relatórios 🛡️</h1>
        <h3 style='color: #60A5FA !important; font-size: 18px; font-weight: 600; margin-top: 0px; margin-bottom: 4px;'>SEÇÃO — PM/3</h3>
        <h4 style='color: #E0E6ED !important; font-size: 16px; font-weight: 500; margin-top: 0px; margin-bottom: 2px;'>18º BPM</h4>
        <h5 style='color: #9CA3AF !important; font-size: 14px; font-weight: normal; margin-top: 0px; margin-bottom: 15px;'>2º CRPM</h5>
    </div>
    """, unsafe_allow_html=True)

    st.markdown("<p style='text-align: center; color: #9CA3AF;'>Selecione abaixo o sistema que deseja acessar:</p>", unsafe_allow_html=True)
    st.write("")

    col_s1, col_s2 = st.columns(2)

    with col_s1:
        st.markdown("""
        <div class="portal-card">
            <h3>📑 Gerador de OS</h3>
            <p>Extração automática de dados de Ordens de Operação e geração padronizada de Ordens de Serviço em Word e PDF.</p>
        </div>
        """, unsafe_allow_html=True)
        if st.button("Acessar Gerador de OS ➔", key="btn_os"):
            st.session_state.pagina_atual = "gerador_os"
            st.rerun()

    with col_s2:
        st.markdown("""
        <div class="portal-card">
            <h3>🗓️ Extrajornada / Escalas</h3>
            <p>Gestão, conferência e montagem da programação de escalas de serviço extrajornada voluntária do 18º BPM.</p>
        </div>
        """, unsafe_allow_html=True)
        if st.button("Acessar Extrajornada ➔", key="btn_extrajornada"):
            st.session_state.pagina_atual = "segundo_site"
            st.rerun()


# =========================================================
# ROTEADOR DE NAVEGAÇÃO CENTRAL
# =========================================================

if st.session_state.pagina_atual != "portal":
    with st.sidebar:
        if caminho_brasao:
            st.image(caminho_brasao, width=90)
        st.write("### 🧭 Navegação")
        if st.button("⬅️ Voltar ao Portal Principal"):
            st.session_state.pagina_atual = "portal"
            st.rerun()
        st.markdown("---")

if st.session_state.pagina_atual == "portal":
    render_tela_entrada()
elif st.session_state.pagina_atual == "gerador_os":
    render_modulo_gerador_os()
elif st.session_state.pagina_atual == "segundo_site":
    render_modulo_segundo_site()
